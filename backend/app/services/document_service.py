"""Document service for generating official confirmation attachments from templates."""
import base64
import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import docx

from app.core.config import get_settings

logger = logging.getLogger(__name__)

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent.parent
    / "doc"
    / "Car_Rental_Payment_Authorization_Confirmation.docx"
)


def _replace_in_paragraph(p: Any, old_text: str, new_text: str) -> None:
    """Safely replaces text within a paragraph while preserving runs/formatting where possible."""
    while old_text in p.text:
        found = False
        for run in p.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text, 1)
                found = True
                break
        if not found:
            full_text = p.text.replace(old_text, new_text, 1)
            if p.runs:
                p.runs[0].text = full_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = full_text
            break


def generate_confirmation_docx(
    *,
    lead_id: str,
    customer_name: str,
    customer_email: str,
    booking_reference: str = "",
    confirmation_number: str = "",
    service_type: str = "car",
    car_provider: str = "",
    booking_platform: str = "",
    vehicle_type: str = "",
    car_model: str = "",
    driver_name: str = "",
    pickup_datetime_str: str = "",
    pickup_location: str = "",
    return_datetime_str: str = "",
    return_location: str = "",
    duration_str: str = "1 Day",
    prepaid_amount: float = 0.0,
    pay_at_counter_amount: float = 0.0,
    total_amount: float = 0.0,
    client_ip: str = "127.0.0.1",
    user_agent: str = "Authorized Portal",
) -> tuple[str, str]:
    """Populates Car_Rental_Payment_Authorization_Confirmation.docx with actual lead & booking data.
    Returns (filename: str, base64_content: str).
    """
    settings = get_settings()
    brand_name = settings.resend_from_name or "E-Booking Desk"
    support_phone = settings.support_phone or "+1 (877) 362-2838"
    support_email = settings.resend_from_email or "support@chaudharytechblog.co.in"
    support_website = "www.chaudharytechblog.co.in"

    now_utc = datetime.now(timezone.utc)
    date_str = now_utc.strftime("%d-%b-%Y")
    time_str = now_utc.strftime("%H:%M:%S UTC")

    ref_display = booking_reference or f"CRM-{lead_id[:6].upper()}"
    conf_display = confirmation_number or ref_display
    lead_driver = driver_name or customer_name or "Lead Driver"
    provider_display = car_provider or "Car Rental"
    platform_display = booking_platform or "Direct"
    v_type = vehicle_type.replace("_", " ").title() if vehicle_type else "Standard"
    v_model = car_model or v_type

    pickup_full = f"{pickup_datetime_str} – {pickup_location}" if pickup_location else (pickup_datetime_str or "Confirmed")
    return_full = f"{return_datetime_str} – {return_location}" if return_location else (return_datetime_str or "Confirmed")

    calc_total = total_amount if total_amount > 0 else (prepaid_amount + pay_at_counter_amount)

    replacements = {
        "[Booking Reference]": ref_display,
        "[Customer Name]": customer_name or "Valued Customer",
        "[Agency Name]": brand_name,
        "[CRM ID]": ref_display,
        "[Booking Number]": conf_display,
        "[Car Provider]": provider_display,
        "[Booking Platform]": platform_display,
        "[Vehicle Category]": v_type,
        "[Vehicle Model]": v_model,
        "[Date & Time] – [Pickup Location]": pickup_full,
        "[Date & Time] – [Return Location]": return_full,
        "[Number of Days]": duration_str,
        "[Prepaid Amount]": f"{prepaid_amount:.2f}",
        "[Pay Later Amount]": f"{pay_at_counter_amount:.2f}",
        "[Total Amount]": f"{calc_total:.2f}",
        "[Rental Company]": provider_display,
        "[Date]": date_str,
        "[Time + Time Zone]": time_str,
        "[Customer Email]": customer_email or "customer@example.com",
        "[IP Address]": client_ip or "127.0.0.1",
        "[Device/Browser]": user_agent or "Authorized System",
        "[Phone Number]": f" {support_phone} ",
        "[Email Address]": f" {support_email} ",
        "[Website]": f" {support_website}",
    }

    if not TEMPLATE_PATH.exists():
        logger.error(f"[Document Service] Template file not found at {TEMPLATE_PATH}")
        raise FileNotFoundError(f"Confirmation template not found at {TEMPLATE_PATH}")

    doc = docx.Document(str(TEMPLATE_PATH))

    # Replace in body paragraphs
    for p in doc.paragraphs:
        for k, v in replacements.items():
            if k in p.text:
                _replace_in_paragraph(p, k, v)

    # Replace in tables
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for k, v in replacements.items():
                        if k in p.text:
                            _replace_in_paragraph(p, k, v)

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()
    b64_content = base64.b64encode(file_bytes).decode("utf-8")

    filename = f"Car_Rental_Payment_Authorization_Confirmation_{ref_display}.docx"
    logger.info(f"[Document Service] Successfully generated confirmation document {filename} ({len(file_bytes)} bytes)")
    return filename, b64_content
